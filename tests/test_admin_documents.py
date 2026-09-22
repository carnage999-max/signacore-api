from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

import fitz
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core import mail
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document as WordDocument
from rest_framework.test import APIClient

from apps.accounts.models import Organization, OrganizationMembership
from apps.billing.models import OrganizationSubscription
from apps.documents.models import AdminAuditLog, Document, DocumentField
from apps.signing.models import SigningRequest


def build_acroform_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)

    text_widget = fitz.Widget()
    text_widget.field_name = "employee_name"
    text_widget.field_label = "Employee Name"
    text_widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    text_widget.rect = fitz.Rect(72, 144, 240, 168)
    text_widget.field_flags = 1 << 1
    page.add_widget(text_widget)

    signature_widget = fitz.Widget()
    signature_widget.field_name = "employee_signature"
    signature_widget.field_label = "Employee Signature"
    signature_widget.field_type = fitz.PDF_WIDGET_TYPE_SIGNATURE
    signature_widget.rect = fitz.Rect(72, 220, 280, 260)
    signature_widget.field_flags = 1 << 1
    page.add_widget(signature_widget)

    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def build_acroform_pdf_with_long_label() -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)

    text_widget = fitz.Widget()
    text_widget.field_name = "employee_name"
    text_widget.field_label = "L" * 300
    text_widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    text_widget.rect = fitz.Rect(72, 144, 240, 168)
    page.add_widget(text_widget)

    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def build_flat_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((72, 120), "Name:")
    page.insert_text((72, 220), "Initials:")
    shape = page.new_shape()
    shape.draw_line((72, 320), (240, 320))
    shape.finish(width=1)
    shape.commit()

    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def build_heuristic_signature_checkbox_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((72, 120), "Tenant Signature:________________________ Date:____________")
    page.insert_text((72, 150), "Print Name:________________________")
    page.insert_text((72, 180), "Email:_____________________________")
    page.insert_text((92, 220), "I agree to receive notices by email.")
    shape = page.new_shape()
    shape.draw_rect(fitz.Rect(72, 210, 82, 220))
    shape.finish(width=1, color=(0, 0, 0))
    shape.commit()

    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def build_docx() -> bytes:
    document = WordDocument()
    document.add_heading("Imported agreement", level=1)
    document.add_paragraph("Review these terms before sending the agreement.")
    document.add_paragraph("First requirement", style="List Bullet")
    document.add_paragraph("Second requirement", style="List Bullet")
    document.add_paragraph("Final approval", style="List Number")
    payload = BytesIO()
    document.save(payload)
    return payload.getvalue()


@override_settings(
    SIGNACORE_SHARED_SECRET="test-signacore-secret",
    SIGNACORE_SERVICE_USERNAME="signacore-service",
    SIGNACORE_APP_URL="https://mysignacore.com",
    SIGNACORE_SIGNER_PORTAL_URL="https://sign.mysignacore.com",
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    CELERY_TASK_ALWAYS_EAGER=True,
    CELERY_TASK_EAGER_PROPAGATES=True,
)
class AdminDocumentUploadTests(TestCase):
    def setUp(self) -> None:
        self.client = APIClient()
        self.user = get_user_model().objects.create_user(
            username="admin",
            email="admin@mysignacore.com",
            password="password123",
            is_staff=True,
        )
        self.organization = Organization.objects.create(name="Test Company", created_by=self.user)
        OrganizationMembership.objects.create(
            organization=self.organization,
            user=self.user,
            role=OrganizationMembership.RoleEnum.ADMIN,
        )
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(self.user.id),
            HTTP_X_SIGNACORE_ORGANIZATION_ID=str(self.organization.id),
        )

    def test_free_plan_blocks_the_sixth_document_this_month(self) -> None:
        for index in range(5):
            Document.objects.create(
                title=f"Free document {index}",
                original_pdf=SimpleUploadedFile(
                    f"free-{index}.pdf",
                    build_flat_pdf(),
                    content_type="application/pdf",
                ),
                created_by=self.user,
                organization=self.organization,
            )

        response = self.client.post(
            "/api/admin/documents/",
            {
                "title": "Sixth document",
                "pdf_file": SimpleUploadedFile("sixth.pdf", build_flat_pdf(), content_type="application/pdf"),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 403, response.json())
        self.assertEqual(response.json()["code"], "FREE_MONTHLY_LIMIT_REACHED")
        self.assertEqual(Document.objects.filter(organization=self.organization).count(), 5)

    def test_professional_subscription_allows_audit_history(self) -> None:
        OrganizationSubscription.objects.create(
            organization=self.organization,
            plan=OrganizationSubscription.PlanEnum.PROFESSIONAL,
            status=OrganizationSubscription.StatusEnum.ACTIVE,
        )

        response = self.client.get("/api/admin/audit-logs/")

        self.assertEqual(response.status_code, 200, response.json())

    def test_free_plan_cannot_view_audit_history(self) -> None:
        response = self.client.get("/api/admin/audit-logs/")

        self.assertEqual(response.status_code, 403, response.json())
        self.assertEqual(response.json()["code"], "PLAN_FEATURE_REQUIRED")

    def test_admin_routes_require_signacore_secret_header(self) -> None:
        self.client.credentials()

        response = self.client.get("/api/admin/documents/")

        self.assertEqual(response.status_code, 403, response.json())
        self.assertEqual(response.json()["detail"], "Invalid Signacore secret.")

        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(self.user.id),
            HTTP_X_SIGNACORE_ORGANIZATION_ID=str(self.organization.id),
        )

    def test_admin_login_is_rate_limited_by_proxy_ip(self) -> None:
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_REAL_IP="198.51.100.221",
        )
        responses = [
            self.client.post(
                "/api/admin/auth/login/",
                {"identifier": "admin", "password": "wrong-password"},
                format="json",
            )
            for _ in range(6)
        ]

        self.assertTrue(all(response.status_code == 400 for response in responses[:5]))
        self.assertEqual(responses[-1].status_code, 429)

    def test_upload_rejects_a_pdf_over_the_configured_size_limit(self) -> None:
        with self.settings(SIGNACORE_MAX_DOCUMENT_UPLOAD_BYTES=1024 * 1024):
            response = self.client.post(
                "/api/admin/documents/",
                {
                    "title": "Oversized PDF",
                    "pdf_file": SimpleUploadedFile(
                        "oversized.pdf",
                        b"%PDF-1.4\n" + (b"0" * (1024 * 1024)),
                        content_type="application/pdf",
                    ),
                },
                format="multipart",
            )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertEqual(response.json()["pdf_file"], ["PDF files must be 1 MB or smaller."])

    def test_upload_pdf_creates_document_and_extracts_acroform_fields(self) -> None:
        upload = SimpleUploadedFile(
            "employment.pdf",
            build_acroform_pdf(),
            content_type="application/pdf",
        )

        response = self.client.post(
            "/api/admin/documents/",
            {"title": "Employment Agreement", "pdf_file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertEqual(payload["title"], "Employment Agreement")
        self.assertEqual(payload["status"], "DRAFT")
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(payload["detection_summary"]["source"], "ACROFORM")
        self.assertEqual(payload["detection_summary"]["field_count"], 2)
        self.assertEqual(payload["detection_summary"]["native_widget_count"], 2)
        self.assertEqual(payload["detection_summary"]["ignored_widget_count"], 0)
        self.assertEqual(len(payload["fields"]), 2)
        self.assertEqual(
            [field["field_type"] for field in payload["fields"]],
            ["TEXT", "SIGNATURE"],
        )
        document = Document.objects.get(pk=payload["id"])
        self.assertEqual(document.created_by, self.user)
        self.assertEqual(document.organization, self.organization)
        self.assertEqual(document.fields.count(), 2)

    def test_authored_document_creates_encrypted_pdf_and_exact_fields(self) -> None:
        response = self.client.post(
            "/api/admin/documents/authored/",
            {
                "title": "Created agreement",
                "content": {
                    "type": "doc",
                    "attrs": {"pageSize": "LETTER"},
                    "content": [
                        {
                            "type": "heading",
                            "attrs": {"level": 1},
                            "content": [{"type": "text", "text": "Created agreement"}],
                        },
                        {
                            "type": "paragraph",
                            "content": [{"type": "text", "text": "Please review these terms."}],
                        },
                        {
                            "type": "signacoreField",
                            "attrs": {
                                "fieldType": "SIGNATURE",
                                "label": "Customer signature",
                                "required": True,
                            },
                        },
                    ],
                },
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        document = Document.objects.get(pk=payload["id"])
        self.assertEqual(document.source, Document.SourceEnum.AUTHORED)
        self.assertTrue(document.original_pdf.name.endswith(".pdf"))
        self.assertTrue(document.authored_content)
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(len(payload["fields"]), 1)
        self.assertEqual(payload["fields"][0]["field_type"], "SIGNATURE")
        self.assertEqual(payload["fields"][0]["detection_source"], "AUTHORED")
        self.assertGreater(payload["fields"][0]["width"], 0)

    def test_docx_import_returns_native_editor_content_without_creating_document(self) -> None:
        response = self.client.post(
            "/api/admin/documents/authored/import/",
            {
                "docx_file": SimpleUploadedFile(
                    "employment-agreement.docx",
                    build_docx(),
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload["title"], "employment-agreement")
        self.assertEqual(payload["content"]["type"], "doc")
        self.assertEqual(
            [node["type"] for node in payload["content"]["content"]],
            ["heading", "paragraph", "bulletList", "orderedList"],
        )
        self.assertEqual(payload["warnings"], [])
        self.assertTrue(
            AdminAuditLog.objects.filter(
                organization=self.organization,
                action=AdminAuditLog.ActionEnum.DOCUMENT_IMPORT,
            ).exists()
        )
        self.assertEqual(Document.objects.filter(organization=self.organization).count(), 0)

    def test_docx_import_rejects_non_docx_files(self) -> None:
        response = self.client.post(
            "/api/admin/documents/authored/import/",
            {"docx_file": SimpleUploadedFile("agreement.pdf", b"not a docx", content_type="application/pdf")},
            format="multipart",
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertIn("docx_file", response.json())

    def test_authored_document_rejects_unknown_content_nodes(self) -> None:
        response = self.client.post(
            "/api/admin/documents/authored/",
            {
                "title": "Invalid agreement",
                "content": {
                    "type": "doc",
                    "content": [{"type": "video", "attrs": {}}],
                },
            },
            format="json",
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertIn("content", response.json())

    def test_authored_document_update_rebuilds_fields_and_pdf(self) -> None:
        create_response = self.client.post(
            "/api/admin/documents/authored/",
            {
                "title": "Draft agreement",
                "content": {
                    "type": "doc",
                    "attrs": {"pageSize": "A4"},
                    "content": [
                        {
                            "type": "signacoreField",
                            "attrs": {"fieldType": "TEXT", "label": "Name", "required": True},
                        }
                    ],
                },
            },
            format="json",
        )
        document_id = create_response.json()["id"]

        update_response = self.client.patch(
            f"/api/admin/documents/{document_id}/authored/",
            {
                "title": "Updated agreement",
                "content": {
                    "type": "doc",
                    "attrs": {"pageSize": "A4"},
                    "content": [
                        {
                            "type": "signacoreField",
                            "attrs": {"fieldType": "INITIALS", "label": "Initials", "required": False},
                        },
                        {
                            "type": "signacoreField",
                            "attrs": {"fieldType": "SIGNATURE", "label": "Signature", "required": True},
                        },
                    ],
                },
            },
            format="json",
        )

        self.assertEqual(update_response.status_code, 200, update_response.json())
        document = Document.objects.get(pk=document_id)
        self.assertEqual(document.title, "Updated agreement")
        self.assertEqual(document.fields.count(), 2)
        self.assertEqual(
            list(document.fields.values_list("field_type", flat=True)),
            ["INITIALS", "SIGNATURE"],
        )

        reopen_response = self.client.get(f"/api/admin/documents/{document_id}/authored/")
        self.assertEqual(reopen_response.status_code, 200, reopen_response.json())
        self.assertEqual(reopen_response.json()["content"]["attrs"]["pageSize"], "A4")

    def test_authored_document_generates_additional_pages_without_losing_fields(self) -> None:
        response = self.client.post(
            "/api/admin/documents/authored/",
            {
                "title": "Long agreement",
                "content": {
                    "type": "doc",
                    "attrs": {"pageSize": "LETTER"},
                    "content": [
                        {
                            "type": "paragraph",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "This paragraph is repeated to verify deterministic page breaks in the authored document renderer.",
                                }
                            ],
                        }
                        for _ in range(60)
                    ]
                    + [
                        {
                            "type": "signacoreField",
                            "attrs": {"fieldType": "SIGNATURE", "label": "Final signature", "required": True},
                        }
                    ],
                },
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertGreater(payload["page_count"], 1)
        self.assertEqual(payload["fields"][0]["page"], payload["page_count"])

    def test_upload_pdf_falls_back_to_heuristic_detection_when_no_widgets_exist(self) -> None:
        upload = SimpleUploadedFile(
            "flat-contract.pdf",
            build_flat_pdf(),
            content_type="application/pdf",
        )

        response = self.client.post(
            "/api/admin/documents/",
            {"title": "Flat Contract", "pdf_file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertEqual(payload["detection_summary"]["source"], "HEURISTIC")
        self.assertGreaterEqual(payload["detection_summary"]["field_count"], 2)
        self.assertTrue(any(field["field_type"] == "SIGNATURE" for field in payload["fields"]))
        self.assertTrue(any(field["field_type"] == "TEXT" for field in payload["fields"]))
        self.assertTrue(
            all(
                8 <= field["height"] <= 16
                for field in payload["fields"]
                if field["field_type"] in {"SIGNATURE", "TEXT"}
            )
        )

    def test_upload_pdf_detects_inline_signature_date_text_and_checkbox_fields(self) -> None:
        upload = SimpleUploadedFile(
            "heuristic-fields.pdf",
            build_heuristic_signature_checkbox_pdf(),
            content_type="application/pdf",
        )

        response = self.client.post(
            "/api/admin/documents/",
            {"title": "Heuristic Fields", "pdf_file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertEqual(payload["detection_summary"]["source"], "HEURISTIC")
        self.assertEqual(len(payload["fields"]), 5)

        returned_fields = {(field["label"], field["field_type"]) for field in payload["fields"]}
        self.assertIn(("Tenant Signature", "SIGNATURE"), returned_fields)
        self.assertIn(("Date", "TEXT"), returned_fields)
        self.assertIn(("Print Name", "TEXT"), returned_fields)
        self.assertIn(("Email", "TEXT"), returned_fields)
        self.assertTrue(
            any(
                field["field_type"] == "CHECKBOX"
                and "I agree to receive notices by email" in field["label"]
                and field["width"] <= 16
                for field in payload["fields"]
            )
        )

    def test_upload_pdf_truncates_overlong_detected_labels(self) -> None:
        upload = SimpleUploadedFile(
            "long-label.pdf",
            build_acroform_pdf_with_long_label(),
            content_type="application/pdf",
        )

        response = self.client.post(
            "/api/admin/documents/",
            {"title": "Long Label Contract", "pdf_file": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertEqual(len(payload["fields"][0]["label"]), 255)

    def test_list_documents_returns_signer_progress_counts(self) -> None:
        first = Document.objects.create(
            title="One",
            original_pdf=SimpleUploadedFile("one.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )
        second = Document.objects.create(
            title="Two",
            original_pdf=SimpleUploadedFile("two.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.SENT,
        )
        DocumentField.objects.create(
            document=second,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Name",
            page=1,
            x=10,
            y=10,
            width=100,
            height=20,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )

        response = self.client.get("/api/admin/documents/")

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(len(payload["items"]), 2)
        returned_ids = {item["id"] for item in payload["items"]}
        self.assertEqual(returned_ids, {str(first.id), str(second.id)})

    def test_document_detail_returns_fields(self) -> None:
        document = Document.objects.create(
            title="Offer Letter",
            original_pdf=SimpleUploadedFile("offer.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )
        field = DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Employee Name",
            page=1,
            x=20,
            y=40,
            width=120,
            height=20,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )

        response = self.client.get(f"/api/admin/documents/{document.id}/")

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload["id"], str(document.id))
        self.assertEqual(len(payload["fields"]), 1)
        self.assertEqual(payload["fields"][0]["id"], str(field.id))
        self.assertEqual(payload["page_count"], 1)
        self.assertEqual(len(payload["pages"]), 1)
        self.assertIn(f"/api/admin/documents/{document.id}/pages/1/preview/", payload["pages"][0]["preview_url"])

    @patch("apps.documents.views.serialize_document_detail", side_effect=OSError("encrypted PDF is unavailable"))
    def test_document_detail_returns_json_when_pdf_preview_generation_fails(self, serialize_detail) -> None:
        document = Document.objects.create(
            title="Unavailable Preview",
            original_pdf=SimpleUploadedFile("unavailable.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        with self.assertLogs("apps.documents.views", level="ERROR") as logs:
            response = self.client.get(f"/api/admin/documents/{document.id}/")

        self.assertEqual(response.status_code, 503, response.json())
        self.assertEqual(
            response.json()["detail"],
            "This document could not be opened because its PDF preview is unavailable. Try again shortly.",
        )
        self.assertTrue(any("Admin document detail serialization failed" in message for message in logs.output))
        serialize_detail.assert_called_once_with(document)

    def test_document_page_preview_returns_png(self) -> None:
        document = Document.objects.create(
            title="Previewable",
            original_pdf=SimpleUploadedFile("preview.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        response = self.client.get(f"/api/admin/documents/{document.id}/pages/1/preview/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "image/png")

    def test_create_manual_field_persists_to_document(self) -> None:
        document = Document.objects.create(
            title="NDA",
            original_pdf=SimpleUploadedFile("nda.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/fields/",
            {
                "field_type": "INITIALS",
                "label": "Employee Initials",
                "page": 1,
                "x": 72,
                "y": 120,
                "width": 48,
                "height": 24,
                "is_required": True,
                "order": 1,
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.json())
        payload = response.json()
        self.assertEqual(payload["field_type"], "INITIALS")
        self.assertEqual(payload["detection_source"], "MANUAL")
        self.assertEqual(document.fields.count(), 1)

    def test_patch_field_updates_coordinates_and_required_flag(self) -> None:
        document = Document.objects.create(
            title="Agreement",
            original_pdf=SimpleUploadedFile("agreement.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )
        field = DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Name",
            page=1,
            x=20,
            y=20,
            width=100,
            height=20,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )

        response = self.client.patch(
            f"/api/admin/documents/{document.id}/fields/{field.id}/",
            {"x": 88, "y": 144, "is_required": False, "label": "Full Legal Name"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        field.refresh_from_db()
        self.assertEqual(field.x, 88)
        self.assertEqual(field.y, 144)
        self.assertFalse(field.is_required)
        self.assertEqual(field.label, "Full Legal Name")

    def test_delete_field_removes_it(self) -> None:
        document = Document.objects.create(
            title="Policy",
            original_pdf=SimpleUploadedFile("policy.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )
        field = DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Date",
            page=1,
            x=40,
            y=80,
            width=90,
            height=20,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )

        response = self.client.delete(f"/api/admin/documents/{document.id}/fields/{field.id}/")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(DocumentField.objects.filter(id=field.id).exists())

    def test_delete_document_removes_document_and_audit_log(self) -> None:
        document = Document.objects.create(
            title="Delete me",
            original_pdf=SimpleUploadedFile("delete.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        response = self.client.delete(f"/api/admin/documents/{document.id}/")

        self.assertEqual(response.status_code, 204)
        self.assertFalse(Document.objects.filter(pk=document.id).exists())
        self.assertTrue(
            AdminAuditLog.objects.filter(
                organization=self.organization,
                action=AdminAuditLog.ActionEnum.DOCUMENT_DELETE,
                target_id=str(document.id),
            ).exists()
        )

    def test_send_document_creates_signing_requests_and_marks_document_sent(self) -> None:
        mail.outbox = []
        document = Document.objects.create(
            title="Offer Package",
            original_pdf=SimpleUploadedFile("offer.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )
        DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.SIGNATURE,
            label="Signature",
            page=1,
            x=72,
            y=120,
            width=180,
            height=36,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/send/",
            {
                "signers": [
                    {"signer_email": "jane@example.com", "signer_name": "Jane Doe"},
                    {"signer_email": "john@example.com", "signer_name": "John Doe"},
                ]
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload["status"], "SENT")
        self.assertEqual(payload["signer_count"], 2)
        self.assertEqual(len(payload["signing_requests"]), 2)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.StatusEnum.SENT)
        self.assertEqual(document.signing_requests.count(), 2)
        self.assertEqual(len(mail.outbox), 2)
        self.assertIn("https://sign.mysignacore.com/sign/", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].alternatives[0][1], "text/html")
        self.assertIn("Review and sign", mail.outbox[0].alternatives[0][0])
        self.assertIn("Offer Package", mail.outbox[0].alternatives[0][0])

    def test_send_document_requires_at_least_one_field(self) -> None:
        document = Document.objects.create(
            title="Blank Contract",
            original_pdf=SimpleUploadedFile("blank.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/send/",
            {"signers": [{"signer_email": "jane@example.com"}]},
            format="json",
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertIn("fields", response.json())

    def test_send_document_allows_adding_signers_to_sent_document(self) -> None:
        document = Document.objects.create(
            title="Existing Sent Doc",
            original_pdf=SimpleUploadedFile("sent.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.SENT,
        )
        DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Name",
            page=1,
            x=72,
            y=120,
            width=180,
            height=24,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )
        SigningRequest.objects.create(
            document=document,
            signer_email="existing@example.com",
            signer_name="Existing User",
            expires_at=timezone.now(),
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/send/",
            {"signers": [{"signer_email": "another@example.com"}]},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        document.refresh_from_db()
        self.assertEqual(document.status, Document.StatusEnum.SENT)
        self.assertEqual(document.signing_requests.count(), 2)

    def test_send_document_rejects_duplicate_active_signer_email(self) -> None:
        document = Document.objects.create(
            title="Duplicate Check",
            original_pdf=SimpleUploadedFile("duplicate.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.SENT,
        )
        DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Name",
            page=1,
            x=72,
            y=120,
            width=180,
            height=24,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )
        SigningRequest.objects.create(
            document=document,
            signer_email="existing@example.com",
            signer_name="Existing User",
            expires_at=timezone.now() + timedelta(days=7),
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/send/",
            {"signers": [{"signer_email": "existing@example.com"}]},
            format="json",
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertIn("signers", response.json())

    def test_patch_document_updates_title(self) -> None:
        document = Document.objects.create(
            title="Old Title",
            original_pdf=SimpleUploadedFile("doc.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
        )

        response = self.client.patch(
            f"/api/admin/documents/{document.id}/",
            {"title": "New Title"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        document.refresh_from_db()
        self.assertEqual(document.title, "New Title")

    def test_void_document_marks_pending_requests_invalid(self) -> None:
        document = Document.objects.create(
            title="To Void",
            original_pdf=SimpleUploadedFile("void.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.SENT,
        )
        request = SigningRequest.objects.create(
            document=document,
            signer_email="void@example.com",
            signer_name="Void User",
            expires_at=timezone.now() + timedelta(days=7),
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/void/",
            {"voided_reason": "Sent in error"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        document.refresh_from_db()
        request.refresh_from_db()
        self.assertEqual(document.status, Document.StatusEnum.VOIDED)
        self.assertEqual(document.voided_reason, "Sent in error")
        self.assertIsNotNone(document.voided_at)
        self.assertLessEqual(request.expires_at, timezone.now())

    def test_resend_signing_request_reopens_signed_document_and_clears_prior_submission(self) -> None:
        document = Document.objects.create(
            title="Reopen Me",
            original_pdf=SimpleUploadedFile("reopen.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.COMPLETED,
        )
        document.signed_pdf.save(
            "reopen-signed.pdf",
            ContentFile(build_flat_pdf()),
            save=True,
        )
        field = DocumentField.objects.create(
            document=document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Name",
            page=1,
            x=72,
            y=120,
            width=180,
            height=24,
            is_required=True,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )
        signing_request = SigningRequest.objects.create(
            document=document,
            signer_email="signed@example.com",
            signer_name="Signed User",
            status=SigningRequest.StatusEnum.SIGNED,
            signed_at=timezone.now(),
            expires_at=timezone.now() + timedelta(days=7),
            ip_address="127.0.0.1",
            user_agent="Browser",
        )
        signing_request.submissions.create(
            document_field=field,
            value_type="TEXT",
            text_value="Signed Value",
        )

        response = self.client.post(
            f"/api/admin/documents/{document.id}/signing-requests/{signing_request.id}/resend/",
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        document.refresh_from_db()
        signing_request.refresh_from_db()
        self.assertEqual(document.status, Document.StatusEnum.SENT)
        self.assertFalse(bool(document.signed_pdf))
        self.assertEqual(signing_request.status, SigningRequest.StatusEnum.PENDING)
        self.assertIsNone(signing_request.signed_at)
        self.assertEqual(signing_request.submissions.count(), 0)

    def test_download_signed_document_returns_file(self) -> None:
        document = Document.objects.create(
            title="Completed",
            original_pdf=SimpleUploadedFile("original.pdf", build_flat_pdf(), content_type="application/pdf"),
            created_by=self.user,
            organization=self.organization,
            status=Document.StatusEnum.COMPLETED,
        )
        document.signed_pdf.save(
            "signed.pdf",
            ContentFile(build_flat_pdf()),
            save=True,
        )

        response = self.client.get(f"/api/admin/documents/{document.id}/download/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

    def test_staff_admin_can_login_with_django_credentials(self) -> None:
        response = self.client.post(
            "/api/admin/auth/login/",
            {"identifier": "admin", "password": "password123"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertEqual(payload["admin"]["username"], "admin")
        self.assertFalse(payload["admin"]["is_superuser"])
        self.assertEqual(payload["admin"]["organizations"][0]["id"], str(self.organization.id))
        self.assertEqual(payload["admin"]["organizations"][0]["role"], "ADMIN")
        self.assertTrue(
            AdminAuditLog.objects.filter(
                action=AdminAuditLog.ActionEnum.LOGIN,
                actor=self.user,
            ).exists()
        )

    def test_superuser_can_create_admin_user_and_email_is_sent(self) -> None:
        superuser = get_user_model().objects.create_superuser(
            username="owner",
            email="owner@mysignacore.com",
            password="password123",
        )
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(superuser.id),
        )

        response = self.client.post(
            "/api/admin/users/",
            {
                "username": "ops-admin",
                "email": "ops@mysignacore.com",
                "first_name": "Ops",
                "last_name": "Admin",
                "password": "TempPass123!",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 201, response.json())
        created_user = get_user_model().objects.get(username="ops-admin")
        self.assertTrue(created_user.is_staff)
        self.assertFalse(created_user.is_superuser)
        self.assertTrue(
            OrganizationMembership.objects.filter(
                organization=self.organization,
                user=created_user,
                role=OrganizationMembership.RoleEnum.ADMIN,
            ).exists()
        )
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Your SignaCore admin account is ready", mail.outbox[0].subject)
        self.assertIn("https://mysignacore.com/admin/login", mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].alternatives[0][1], "text/html")
        self.assertIn("Admin account created", mail.outbox[0].alternatives[0][0])
        self.assertIn("signa-core.png", mail.outbox[0].alternatives[0][0])
        self.assertTrue(
            AdminAuditLog.objects.filter(
                action=AdminAuditLog.ActionEnum.ADMIN_USER_CREATE,
                actor=superuser,
                target_id=str(created_user.id),
            ).exists()
        )

    def test_non_superuser_cannot_create_admin_user(self) -> None:
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(self.user.id),
        )

        response = self.client.post(
            "/api/admin/users/",
            {
                "username": "blocked-admin",
                "email": "blocked@mysignacore.com",
                "password": "TempPass123!",
            },
            format="json",
        )

        self.assertEqual(response.status_code, 403, response.json())
        self.assertFalse(get_user_model().objects.filter(username="blocked-admin").exists())

    def test_superuser_can_change_admin_password(self) -> None:
        superuser = get_user_model().objects.create_superuser(
            username="owner",
            email="owner@mysignacore.com",
            password="password123",
        )
        target = get_user_model().objects.create_user(
            username="reset-me",
            email="reset@mysignacore.com",
            password="OldPass123!",
            is_staff=True,
        )
        OrganizationMembership.objects.create(
            organization=self.organization,
            user=target,
            role=OrganizationMembership.RoleEnum.ADMIN,
        )
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(superuser.id),
        )

        response = self.client.post(
            f"/api/admin/users/{target.id}/password/",
            {"password": "NewTempPass123!"},
            format="json",
        )

        self.assertEqual(response.status_code, 200, response.json())
        target.refresh_from_db()
        self.assertTrue(target.check_password("NewTempPass123!"))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Your SignaCore admin password was changed", mail.outbox[0].subject)
        self.assertEqual(mail.outbox[0].alternatives[0][1], "text/html")
        self.assertIn("Admin password changed", mail.outbox[0].alternatives[0][0])
        self.assertTrue(
            AdminAuditLog.objects.filter(
                action=AdminAuditLog.ActionEnum.ADMIN_PASSWORD_CHANGE,
                actor=superuser,
                target_id=str(target.id),
            ).exists()
        )

    def test_superuser_can_view_audit_logs(self) -> None:
        superuser = get_user_model().objects.create_superuser(
            username="owner",
            email="owner@mysignacore.com",
            password="password123",
        )
        AdminAuditLog.objects.create(
            actor=superuser,
            actor_email=superuser.email,
            action=AdminAuditLog.ActionEnum.LOGIN,
            target_type="admin_user",
            target_id=str(superuser.id),
            summary="Signed in as owner.",
        )
        self.client.credentials(
            HTTP_X_SIGNACORE_SECRET=settings.SIGNACORE_SHARED_SECRET,
            HTTP_X_SIGNACORE_ADMIN_ID=str(superuser.id),
        )

        response = self.client.get("/api/admin/audit-logs/")

        self.assertEqual(response.status_code, 200, response.json())
        self.assertGreaterEqual(len(response.json()["items"]), 1)

    def test_company_members_cannot_see_or_open_another_company_document(self) -> None:
        other_owner = get_user_model().objects.create_user(
            username="other-owner",
            password="password123",
            is_staff=True,
        )
        other_organization = Organization.objects.create(
            name="Other Company",
            created_by=other_owner,
        )
        OrganizationMembership.objects.create(
            organization=other_organization,
            user=other_owner,
            role=OrganizationMembership.RoleEnum.OWNER,
        )
        other_document = Document.objects.create(
            title="Private Other Company Agreement",
            original_pdf=SimpleUploadedFile(
                "other.pdf",
                build_flat_pdf(),
                content_type="application/pdf",
            ),
            created_by=other_owner,
            organization=other_organization,
        )

        list_response = self.client.get("/api/admin/documents/")
        detail_response = self.client.get(f"/api/admin/documents/{other_document.id}/")

        self.assertEqual(list_response.status_code, 200, list_response.json())
        self.assertNotIn(
            str(other_document.id),
            {item["id"] for item in list_response.json()["items"]},
        )
        self.assertEqual(detail_response.status_code, 404, detail_response.json())

    def test_company_members_cannot_mutate_or_download_another_company_document(self) -> None:
        other_owner = get_user_model().objects.create_user(
            username="second-owner",
            password="password123",
            is_staff=True,
        )
        other_organization = Organization.objects.create(
            name="Second Company",
            created_by=other_owner,
        )
        OrganizationMembership.objects.create(
            organization=other_organization,
            user=other_owner,
            role=OrganizationMembership.RoleEnum.OWNER,
        )
        other_document = Document.objects.create(
            title="Second Private Agreement",
            original_pdf=SimpleUploadedFile(
                "second.pdf",
                build_flat_pdf(),
                content_type="application/pdf",
            ),
            created_by=other_owner,
            organization=other_organization,
            status=Document.StatusEnum.SENT,
        )
        other_field = DocumentField.objects.create(
            document=other_document,
            field_type=DocumentField.FieldTypeEnum.TEXT,
            label="Private field",
            page=1,
            x=72,
            y=120,
            width=180,
            height=24,
            detection_source=DocumentField.DetectionSourceEnum.MANUAL,
            order=1,
        )
        other_signing_request = SigningRequest.objects.create(
            document=other_document,
            signer_email="second-signer@example.com",
            expires_at=timezone.now() + timedelta(days=7),
        )

        attempts = [
            self.client.get(f"/api/admin/documents/{other_document.id}/pages/1/preview/"),
            self.client.patch(
                f"/api/admin/documents/{other_document.id}/",
                {"title": "Changed by attacker"},
                format="json",
            ),
            self.client.patch(
                f"/api/admin/documents/{other_document.id}/fields/{other_field.id}/",
                {"label": "Changed by attacker"},
                format="json",
            ),
            self.client.delete(
                f"/api/admin/documents/{other_document.id}/fields/{other_field.id}/",
            ),
            self.client.post(f"/api/admin/documents/{other_document.id}/void/", {}, format="json"),
            self.client.post(
                f"/api/admin/documents/{other_document.id}/send/",
                {"signers": [{"signer_email": "attacker@example.com"}]},
                format="json",
            ),
            self.client.post(
                f"/api/admin/documents/{other_document.id}/signing-requests/{other_signing_request.id}/resend/",
                {},
                format="json",
            ),
            self.client.get(f"/api/admin/documents/{other_document.id}/download/"),
        ]

        self.assertTrue(all(response.status_code == 404 for response in attempts))
        other_document.refresh_from_db()
        other_field.refresh_from_db()
        self.assertEqual(other_document.title, "Second Private Agreement")
        self.assertEqual(other_document.status, Document.StatusEnum.SENT)
        self.assertEqual(other_field.label, "Private field")
        self.assertEqual(other_document.signing_requests.count(), 1)
