from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document as WordDocument

MAX_DOCX_BYTES = 10 * 1024 * 1024
MAX_IMPORTED_PARAGRAPHS = 300


@dataclass(frozen=True)
class DocxImportResult:
    title: str
    content: dict[str, Any]
    warnings: list[str]


class DocxImportError(ValueError):
    pass


class DocxImporter:
    """Convert safe, text-oriented DOCX content into the native editor format."""

    def import_bytes(self, payload: bytes, filename: str) -> DocxImportResult:
        if len(payload) > MAX_DOCX_BYTES:
            raise DocxImportError("DOCX files must be 10 MB or smaller.")
        if not filename.lower().endswith(".docx"):
            raise DocxImportError("Only DOCX files can be imported into the document editor.")
        if not zipfile.is_zipfile(BytesIO(payload)):
            raise DocxImportError("This DOCX file is not valid or could not be opened.")

        try:
            document = WordDocument(BytesIO(payload))
        except Exception as exc:  # python-docx exposes several parser-specific exceptions.
            raise DocxImportError("This DOCX file is not valid or could not be opened.") from exc

        paragraphs = list(document.paragraphs)
        if len(paragraphs) > MAX_IMPORTED_PARAGRAPHS:
            raise DocxImportError("This DOCX contains too many paragraphs for the editor.")

        warnings: list[str] = []
        nodes: list[dict[str, Any]] = []
        list_type: str | None = None
        list_items: list[dict[str, Any]] = []

        def flush_list() -> None:
            nonlocal list_type, list_items
            if list_type and list_items:
                nodes.append({"type": list_type, "content": list_items})
            list_type = None
            list_items = []

        for paragraph in paragraphs:
            text_nodes = self._text_nodes(paragraph)
            text = "".join(str(node.get("text", "")) for node in text_nodes)
            style_name = str(paragraph.style.name or "")
            heading_level = self._heading_level(style_name)
            current_list_type = self._list_type(style_name)

            if current_list_type:
                if list_type != current_list_type:
                    flush_list()
                    list_type = current_list_type
                if text.strip():
                    list_items.append({"type": "listItem", "content": [{"type": "paragraph", "content": text_nodes}]})
                continue

            flush_list()
            if not text.strip():
                nodes.append({"type": "paragraph"})
            elif heading_level:
                nodes.append(
                    {
                        "type": "heading",
                        "attrs": {"level": heading_level},
                        "content": text_nodes,
                    }
                )
            else:
                nodes.append({"type": "paragraph", "content": text_nodes})

        flush_list()

        if not nodes:
            raise DocxImportError("This DOCX does not contain editable text.")
        if document.tables:
            warnings.append("Tables were not imported. Copy their text into the editor and review the layout.")
        if document.inline_shapes:
            warnings.append("Images were not imported. Add them separately before sending the document.")

        return DocxImportResult(
            title=self._title(filename),
            content={"type": "doc", "attrs": {"pageSize": "LETTER"}, "content": nodes},
            warnings=warnings,
        )

    @staticmethod
    def _title(filename: str) -> str:
        title = Path(filename).stem.strip() or "Imported agreement"
        return title[:255]

    @staticmethod
    def _heading_level(style_name: str) -> int | None:
        match = re.search(r"heading\s*([1-3])", style_name, re.IGNORECASE)
        return int(match.group(1)) if match else None

    @staticmethod
    def _list_type(style_name: str) -> str | None:
        normalized = style_name.casefold()
        if "bullet" in normalized:
            return "bulletList"
        if "number" in normalized:
            return "orderedList"
        return None

    @staticmethod
    def _text_nodes(paragraph) -> list[dict[str, Any]]:
        nodes: list[dict[str, Any]] = []
        for run in paragraph.runs:
            if not run.text:
                continue
            node: dict[str, Any] = {"type": "text", "text": run.text}
            marks: list[dict[str, str]] = []
            if run.bold:
                marks.append({"type": "bold"})
            if run.italic:
                marks.append({"type": "italic"})
            if run.underline:
                marks.append({"type": "underline"})
            if marks:
                node["marks"] = marks
            nodes.append(node)
        return nodes or [{"type": "text", "text": paragraph.text}]
